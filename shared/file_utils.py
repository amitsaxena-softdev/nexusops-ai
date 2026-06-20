import io
from pathlib import Path
from docx import Document
import openpyxl


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"=== Sheet: {sheet} ===")
        for row in ws.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def read_sample(path: str) -> bytes:
    return Path(path).read_bytes()


def mime_for(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    return {
        "pdf": "application/pdf",
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
    }.get(ext, "application/octet-stream")


def is_native_gemini(filename: str) -> bool:
    """True for formats Gemini can read natively (PDF, images)."""
    ext = filename.rsplit(".", 1)[-1].lower()
    return ext in {"pdf", "png", "jpg", "jpeg"}
