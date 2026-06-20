import io
import re
from pathlib import Path

import streamlit as st

from agents.shift_agent import find_staff
from shared.file_utils import extract_text_from_xlsx, read_sample
from shared.theme import apply_theme

st.set_page_config(page_title="Shift Replacement — UKS", page_icon="🏥")
apply_theme()
st.title("🏥 Shift Replacement Agent")
st.caption("Client: Universitätsklinikum des Saarlandes (Homburg) — Fill last-minute shift gaps instantly")

SCHEDULE_PATH = Path("samples/hospital_schedule.xlsx")

# ── Schedule source ────────────────────────────────────────────────────────────
uploaded_sched = st.file_uploader(
    "Upload hospital schedule (XLSX) — or leave blank to use sample",
    type=["xlsx"],
    key="sched_upload",
)

if uploaded_sched:
    schedule_bytes = uploaded_sched.read()
    st.success(f"Using uploaded schedule: **{uploaded_sched.name}**")
elif SCHEDULE_PATH.exists():
    schedule_bytes = read_sample(SCHEDULE_PATH)
    st.info("Using sample schedule — 100 UKS staff (hospital_schedule.xlsx)")
else:
    schedule_bytes = None
    st.warning("No schedule available. Upload an XLSX file above.")

schedule_text = extract_text_from_xlsx(schedule_bytes) if schedule_bytes else ""


# ── Structured schedule preview ────────────────────────────────────────────────
def _xlsx_to_dicts(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return [], []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    data = []
    for row in rows[1:]:
        if any(c is not None for c in row):
            record = {}
            for h, c in zip(headers, row):
                if hasattr(c, "strftime"):
                    c = c.strftime("%Y-%m-%d %H:%M")
                record[h] = "" if c is None else c
            data.append(record)
    return headers, data


if schedule_bytes:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(schedule_bytes), data_only=True)

    with st.expander("📅 View roster & weekly schedule"):
        # Shift legend from Shift_Reference sheet
        _, ref_rows = _xlsx_to_dicts(wb, "Shift_Reference")
        ref_codes = [r for r in ref_rows if r.get("Code") and r["Code"] != "Note:"]
        if ref_codes:
            legend = " &nbsp;·&nbsp; ".join(
                f"**`{r['Code']}`** {r['Shift']} {r['Start']}–{r['End']}"
                for r in ref_codes
            )
            st.markdown(f"**Shift codes:** {legend}", unsafe_allow_html=True)

        tab_roster, tab_schedule = st.tabs(["👥 Roster", "📆 Weekly Schedule"])

        with tab_roster:
            _, roster = _xlsx_to_dicts(wb, "Roster")
            # Show key operational columns — hide persona notes, clock times
            show_cols = [
                "Employee ID", "First Name", "Last Name", "Role",
                "Department", "Certifications", "Contract",
                "Max Hrs/Week", "Shift Preference", "Overtime OK", "Status",
            ]
            display = [{k: r.get(k, "") for k in show_cols} for r in roster]
            st.dataframe(display, use_container_width=True, hide_index=True)
            st.caption(f"{len(roster)} staff members · full data (incl. phone, hours) passed to agent")

        with tab_schedule:
            _, sched = _xlsx_to_dicts(wb, "Weekly_Schedule")
            st.dataframe(sched, use_container_width=True, hide_index=True)
            st.caption("O = Off  ·  D = Day (07:00–19:00)  ·  N = Night (19:00–07:00)")

st.divider()

# ── Shift gap input ────────────────────────────────────────────────────────────
DEMO_MESSAGE = (
    "Felix Haddad (HOSP-1059) just called in sick. He was scheduled for tonight's ICU night shift "
    "(19:00–07:00, Saturday 20 June). We need a Registered Nurse with BLS and ACLS. Who can cover?"
)

if "shift_msg" not in st.session_state:
    st.session_state.shift_msg = ""
if "sent_messages" not in st.session_state:
    st.session_state.sent_messages = set()

if st.button("⬇ Load tonight's scenario"):
    st.session_state.shift_msg = DEMO_MESSAGE
    st.session_state.sent_messages = set()
    st.rerun()

message = st.text_area(
    "Message the agent — describe the gap in plain language",
    key="shift_msg",
    placeholder=(
        "e.g. Anna Weber called in sick for tonight's ICU night shift (19:00–07:00). "
        "Need a Registered Nurse with BLS and ACLS. Who's available?"
    ),
    height=110,
)

if st.button("Find Staff & Draft Messages", type="primary") and message:
    st.session_state.sent_messages = set()
    with st.spinner("Checking roster and schedule…"):
        st.session_state.shift_result = find_staff(message, schedule_text)
    st.rerun()


# ── Result rendering ───────────────────────────────────────────────────────────
def _parse_candidates(text):
    DIVIDER = "────────────────────────────────────────"
    parts = text.split(DIVIDER)
    header = parts[0].strip() if parts else ""
    candidates = []
    i = 1
    while i < len(parts):
        block = parts[i].strip()
        if block and "·" in block and "Send:" not in block:
            lines = [l.strip() for l in block.splitlines() if l.strip()]
            name_phone = lines[0] if lines else ""
            role_info  = lines[1] if len(lines) > 1 else ""
            name, phone = "", ""
            if "·" in name_phone:
                np = [x.strip() for x in name_phone.split("·")]
                name  = np[0]
                phone = np[1] if len(np) > 1 else ""
            send_msg = ""
            if i + 1 < len(parts) and "Send:" in parts[i + 1]:
                raw = parts[i + 1].strip()
                m = re.search(r'Send:\s*["“](.+)["”]', raw, re.DOTALL)
                send_msg = m.group(1).strip() if m else raw.replace("Send:", "").strip().strip('"')
                i += 1
            candidates.append({"name": name, "phone": phone,
                                "role_info": role_info, "message": send_msg})
        i += 1
    return header, candidates


def _render_candidates(candidates):
    for c in candidates:
        name     = c["name"]
        phone    = c["phone"]
        role_info = c["role_info"]
        msg      = c["message"]
        sent     = name in st.session_state.sent_messages

        border = "#16a34a" if sent else "#e5e7eb"
        st.markdown(f"""
<div style="border:2px solid {border}; border-radius:10px; padding:16px 20px;
            margin-bottom:14px; background:#f9fafb;">
  <div style="font-size:17px; font-weight:700; color:#111827;">{name}</div>
  <div style="font-size:13px; color:#6b7280; margin-bottom:10px;">{phone} &nbsp;·&nbsp; {role_info}</div>
  <div style="background:#f3f4f6; border-radius:6px; padding:10px 14px;
              font-size:13px; color:#374151; font-family:monospace; white-space:pre-wrap;">{msg}</div>
</div>""", unsafe_allow_html=True)

        if sent:
            st.success(f"✅ WhatsApp sent to {name}")
        else:
            if st.button(f"📤 Send WhatsApp to {name}", key=f"send_{name}"):
                st.session_state.sent_messages.add(name)
                st.rerun()


if "shift_result" in st.session_state:
    result = st.session_state.shift_result
    header, candidates = _parse_candidates(result)

    if candidates:
        st.markdown(f"### {header}" if header else "### Available Staff")
        _render_candidates(candidates)

        all_sent = all(c["name"] in st.session_state.sent_messages for c in candidates)
        if len(candidates) > 1 and not all_sent:
            if st.button("📤 Send WhatsApp to All", type="primary"):
                for c in candidates:
                    st.session_state.sent_messages.add(c["name"])
                st.rerun()
        elif all_sent:
            st.success("✅ All messages sent — shift gap covered.")
    else:
        st.markdown("### Available Staff")
        st.markdown(result)
